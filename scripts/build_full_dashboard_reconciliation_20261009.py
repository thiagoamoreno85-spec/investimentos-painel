from __future__ import annotations

import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

ROOT = Path('/home/ubuntu/investimentos-painel')
XLSX = Path('/home/ubuntu/upload/CONTROLEINVESTIMENTOS-TAM.xlsx')
ASSETS_JSON = Path('/tmp/painel_assets_20261009.json')
OUT = ROOT / 'research' / 'plano_conciliacao_integral_20261009.json'
FX = Decimal('4.9873')
TARGET = Decimal('2075392.00')
PRICE_SCALE = Decimal('0.00000001')
MONEY_SCALE = Decimal('0.01')

ROW_TO_TICKER = {
    15: 'SUZB3', 16: 'KLBN11', 17: 'BRAV3', 18: 'CMIN3', 19: 'VALE3',
    20: 'BPAC11', 21: 'INBR32', 22: 'BBAS3', 23: 'BBDC4', 24: 'CXSE3',
    25: 'AXIA3', 26: 'AXIA7', 27: 'AURE3', 28: 'KEPL3', 29: 'TTEN3',
    30: 'AGRO3', 31: 'SOJA3', 32: 'MBRF3', 33: 'FLRY3', 34: 'SBSP3',
    35: 'ORVR3', 36: 'CYRE3', 37: 'CYRE4', 38: 'ZAVI11', 39: 'XPML11',
    41: 'ELET3', 42: 'BNP_RUBI', 43: 'AZ_QUEST', 44: 'KINEA_GAMA',
    45: 'TREND_INV', 46: 'TC_COSMOS', 47: 'SISPRIME', 48: 'FGTS',
    49: 'NET', 50: 'MSFT', 51: 'GOOGL', 52: 'MU', 53: 'TSM', 54: 'AAPL',
    55: 'NVIDIA', 56: 'CRWV', 57: 'SOUN', 58: 'PGEN', 59: 'TWST', 60: 'UNH',
    61: 'ABBV', 62: 'NEE', 63: 'ENPH', 64: 'PLTR', 65: 'LMT', 66: 'CRWD',
    67: 'TSLA', 68: 'SPCX', 69: 'AMZN', 70: 'CVX', 71: 'BAC', 72: 'BANC',
    73: 'INDA', 74: 'HDB', 75: 'IBN', 76: 'INFY', 77: 'URNM', 78: 'URA',
    79: 'BTC', 80: 'BTC_BIN', 81: 'ETH', 82: 'ETH_BIN', 83: 'SOL',
    84: 'SOL_BIN', 85: 'AVAX', 86: 'BNB', 87: 'AURY', 88: 'FTT',
    91: 'CDB_FIBRA_0528', 92: 'CDB_C6_0227', 93: 'CDB_BMG_0627',
    95: 'CDB_BMG_0127', 96: 'CDB_AGIBANK_0128', 97: 'CRI_MRV_FLEX_0333',
    98: 'CDB_FIBRA_0732', 99: 'CDB_FIBRA_0929', 100: 'NTN_B_052030',
    101: 'CDB_C6_1229', 102: 'CDB_FIBRA_0132', 103: 'DEB_RAIZEN_0630',
    105: 'CRA_MINERVA_0730', 106: 'TESOURO_SELIC_31', 107: 'NTN_B_082032',
    108: 'TESOURO_PRE_0129', 109: 'CDB_XP_0828', 110: 'CDB_C6_IPCA_084_0830',
    111: 'CDB_C6_CONSIG_0229', 112: 'CDB_XP_INV',
    113: 'CDB_PINE_IPCA_0825_0931', 114: 'CRA_MARFRIG_IPCA_1090_0731',
}
KEEP_PANEL_QUANTITY = {
    'TESOURO_SELIC_31', 'NTN_B_082032', 'TESOURO_PRE_0129',
    'CDB_PINE_IPCA_0825_0931', 'CRA_MARFRIG_IPCA_1090_0731',
}


def decimal(value: object, fallback: Decimal | None = None) -> Decimal | None:
    try:
        if value is None:
            return fallback
        return Decimal(str(value))
    except Exception:
        return fallback


wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ws = wb['AÇÕES ']
assets = json.loads(ASSETS_JSON.read_text())
active = {item['ticker']: item for item in assets if decimal(item['totalQuantity'], Decimal('0')) > 0}
plan = []
for row, ticker in ROW_TO_TICKER.items():
    sheet_qty = decimal(ws.cell(row, 8).value, Decimal('0'))
    sheet_position_brl = decimal(ws.cell(row, 14).value, Decimal('0'))
    if ticker not in active:
        if sheet_qty == 0 and sheet_position_brl == 0:
            continue
        raise RuntimeError(f'Linha {row} ({ticker}) possui valor na planilha, mas não existe no painel.')
    before = active[ticker]
    quantity_before = decimal(before['totalQuantity'])
    quantity_after = quantity_before if ticker in KEEP_PANEL_QUANTITY else sheet_qty
    if quantity_after <= 0:
        continue
    currency_after = 'BRL' if ticker == 'BNB' else before['currency']
    unit_price = (sheet_position_brl / quantity_after / (FX if currency_after == 'USD' else Decimal('1'))).quantize(PRICE_SCALE, rounding=ROUND_HALF_UP)
    sheet_cost = decimal(ws.cell(row, 10).value)
    average_cost = sheet_cost if sheet_cost is not None and sheet_cost > 0 else decimal(before['averageCost'])
    total_cost = (quantity_after * average_cost).quantize(MONEY_SCALE, rounding=ROUND_HALF_UP)
    value_brl = quantity_after * unit_price * (FX if currency_after == 'USD' else Decimal('1'))
    plan.append({
        'row': row,
        'ticker': ticker,
        'id': None,
        'name': before['name'],
        'assetClass': before['assetClass'],
        'currencyBefore': before['currency'],
        'currencyAfter': currency_after,
        'quantityBefore': quantity_before,
        'quantityAfter': quantity_after,
        'averageCostBefore': decimal(before['averageCost']),
        'averageCostAfter': average_cost,
        'totalCostBefore': decimal(before['totalCost']),
        'totalCostAfter': total_cost,
        'lastPriceBefore': decimal(before['lastPrice']),
        'lastPriceAfter': unit_price,
        'valueBRLAfter': value_brl,
    })

# A linha 40 representa opção de venda BOVA11; a planilha revisada informa a posição
# e o custo, mas ela ainda não existe na tabela de ativos do painel.
bova_position = decimal(ws.cell(40, 14).value)
bova_cost = decimal(ws.cell(40, 10).value)
plan.append({
    'row': 40,
    'ticker': 'BOVA11_PUT',
    'id': None,
    'name': 'Opção de venda BOVA11',
    'assetClass': 'rv_nacional',
    'currencyBefore': None,
    'currencyAfter': 'BRL',
    'quantityBefore': Decimal('0'),
    'quantityAfter': Decimal('1'),
    'averageCostBefore': Decimal('0'),
    'averageCostAfter': bova_cost,
    'totalCostBefore': Decimal('0'),
    'totalCostAfter': bova_cost,
    'lastPriceBefore': Decimal('0'),
    'lastPriceAfter': bova_position,
    'valueBRLAfter': bova_position,
})

investment_value = sum((item['valueBRLAfter'] for item in plan), Decimal('0'))
cash_after = (TARGET - investment_value).quantize(MONEY_SCALE, rounding=ROUND_HALF_UP)
reconstructed_total = (investment_value + cash_after).quantize(MONEY_SCALE, rounding=ROUND_HALF_UP)

payload = {
    'referenceDate': '2026-10-09',
    'usdBrl': FX,
    'targetTotalBRL': TARGET,
    'assetsMapped': len(plan),
    'investmentValueBRL': investment_value,
    'cashBalanceRequiredBRL': cash_after,
    'reconstructedTotalBRL': reconstructed_total,
    'activePanelTickersNotMapped': sorted(set(active) - {item['ticker'] for item in plan}),
    'updates': plan,
}
OUT.write_text(json.dumps(payload, default=str, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(payload, default=str, ensure_ascii=False, indent=2))
wb.close()
